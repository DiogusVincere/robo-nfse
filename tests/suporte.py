# -*- coding: utf-8 -*-
"""Apoio aos testes: importa o módulo da raiz do repositório e fabrica notas, empresas e XMLs sintéticos.

Tudo aqui é fictício. Os CNPJs saem de `cnpj_ficticio` (dígitos verificadores válidos, empresas inexistentes)."""
import dataclasses
import datetime as dt
import logging
import os
import random
import sys
from pathlib import Path

RAIZ = Path(os.path.abspath(__file__)).parent.parent
if str(RAIZ) not in sys.path:
    sys.path.insert(0, str(RAIZ))

import robo_nfse as robo  # noqa: E402

# sem isto o Python imprime os WARNING do robô (sessão caiu etc.) no meio da saída dos testes
robo.log.addHandler(logging.NullHandler())

MUNICIPIO = '3509502'                 # município do portal de escrituração (padrão do Config)
OUTRO_MUNICIPIO = '3550308'           # um prestador de fora
COMP = robo.Competencia(2026, 9)

_rnd = random.Random(2026)
CNPJ_A = robo.cnpj_ficticio(_rnd)     # 15881399000134 — os mesmos do exemplos/empresas.exemplo.json
CNPJ_B = robo.cnpj_ficticio(_rnd)     # 86987973000181
CNPJ_PRESTADOR = robo.cnpj_ficticio(_rnd)


def empresa(codigo='100', municipio=MUNICIPIO, cnpj=CNPJ_A, responsavel='Analista 1') -> robo.Empresa:
    return robo.Empresa(codigo, f'Empresa Exemplo {codigo} Ltda', cnpj, municipio, responsavel)


def nota(numero='4521', fonte='Nacional', valor=1000.0, aliquota=2.0, municipio=OUTRO_MUNICIPIO,
         prestador=CNPJ_PRESTADOR, tomador=CNPJ_A, retido=False, cancelada=False, comp=COMP, dia=10,
         chave=None, nome='Prestador Exemplo 1', papel='tomada', tomador_nome='', retencoes=None) -> robo.Nota:
    """Uma nota TOMADA pela empresa A (padrão) ou, com papel='prestada', uma nota que a empresa A prestou."""
    if papel == 'prestada':
        prestador, nome, tomador = CNPJ_A, 'Empresa Exemplo 100 Ltda', CNPJ_B
        tomador_nome = tomador_nome or ''
    return robo.Nota(fonte=fonte, numero=numero, chave=chave or f'CHAVE-{fonte}-{numero}',
                     emissao=dt.date(comp.ano, comp.mes, dia), competencia=comp,
                     prestador_cnpj=prestador, prestador_nome=nome, prestador_municipio_ibge=municipio,
                     tomador_cnpj=tomador, tomador_nome=tomador_nome, valor_servico=valor, aliquota=aliquota,
                     valor_iss=round(valor * aliquota / 100, 2), iss_retido=retido, cancelada=cancelada,
                     papel=papel, retencoes=retencoes if retencoes is not None else robo.Retencoes(),
                     xml=f'<NFSe><nNFSe>{numero}</nNFSe><teste>sintética</teste></NFSe>')


def notas_fora(n: int, **kw) -> list:
    """n notas de prestadores de fora do município, numeradas 1001, 1002, ..."""
    return [nota(numero=str(1001 + i), chave=f'CHAVE-{1001 + i}', **kw) for i in range(n)]


def cnpj_valido(cnpj: str) -> bool:
    """Confere os dois dígitos verificadores — é o que o `cnpj_ficticio` promete."""
    d = [int(c) for c in cnpj]
    if len(d) != 14:
        return False
    for tam, pesos in ((12, [5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]), (13, [6, 5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2])):
        resto = sum(x * p for x, p in zip(d[:tam], pesos)) % 11
        if d[tam] != (0 if resto < 2 else 11 - resto):
            return False
    return True


class SemQueda:
    """No lugar do `random.Random` do PortalSimulado: nunca derruba a sessão (random() sempre 1.0)."""

    def random(self):
        return 1.0


def ler_aba(wb, nome):
    linhas = list(wb[nome].iter_rows(values_only=True))
    cab = list(linhas[0])
    return [dict(zip(cab, linha)) for linha in linhas[1:]]


def ler_relatorio(caminho: Path):
    """Lê o Excel do robô: {código: linha da aba Empresas} e a lista de linhas da aba Notas (sem cabeçalho)."""
    from openpyxl import load_workbook
    wb = load_workbook(caminho, read_only=True)
    empresas = {str(linha['Código']): linha for linha in ler_aba(wb, 'Empresas')}
    notas = ler_aba(wb, 'Notas')
    wb.close()
    return empresas, notas


def ler_retencoes(caminho: Path):
    """A aba Retenções: uma linha por empresa × papel, com a soma de cada tributo."""
    from openpyxl import load_workbook
    wb = load_workbook(caminho, read_only=True)
    linhas = ler_aba(wb, 'Retenções')
    wb.close()
    return linhas


def ler_relatorios(pasta: Path):
    """Junta todos os relatórios de `<saida>/_relatorios`."""
    empresas, notas = {}, []
    for arq in sorted((pasta / '_relatorios').glob('*.xlsx')):
        e, n = ler_relatorio(arq)
        empresas.update(e)
        notas += n
    return empresas, notas


__all__ = ['robo', 'dataclasses', 'dt', 'MUNICIPIO', 'OUTRO_MUNICIPIO', 'COMP', 'CNPJ_A', 'CNPJ_B', 'CNPJ_PRESTADOR',
           'empresa', 'nota', 'notas_fora', 'cnpj_valido', 'SemQueda', 'ler_relatorio', 'ler_retencoes',
           'ler_relatorios', 'RAIZ']
