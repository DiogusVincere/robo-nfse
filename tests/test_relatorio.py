# -*- coding: utf-8 -*-
"""Seção 12 (relatório e e-mail), seção 8 (espelho em PDF) e o Progresso da seção 13."""
import importlib.util
import os
import tempfile
import unittest
from pathlib import Path
from unittest import mock

try:
    from . import suporte
except ImportError:
    import suporte

robo = suporte.robo
nota = suporte.nota
COMP = suporte.COMP

CABECALHO_EMPRESAS = ['Código', 'Empresa', 'CNPJ', 'Extraídas', 'Tomadas', 'Prestadas', 'Fora do município',
                      'Escrituradas agora', 'Já escrituradas', 'Guia ISS', 'Retido nas tomadas', 'Retido nas prestadas',
                      'Pendências']
CABECALHO_NOTAS = ['Código', 'Papel', 'Fonte', 'Nº', 'Emissão', 'Prestador', 'CNPJ prestador', 'Tomador', 'CNPJ tomador',
                   'Município (IBGE)', 'Valor', 'Alíquota %', 'ISS', 'ISS retido', 'INSS retido', 'IRRF retido',
                   'PIS retido', 'COFINS retido', 'CSLL retida', 'CSRF retida', 'Total retido', 'Valor líquido',
                   'Retenções', 'Situação']
CABECALHO_RETENCOES = ['Código', 'Empresa', 'Papel', 'Notas com retenção', 'ISS', 'INSS', 'IRRF', 'PIS', 'COFINS', 'CSLL',
                       'CSRF (PIS+COFINS+CSLL)', 'Total']


class TestGerarRelatorio(unittest.TestCase):
    def test_tres_abas_com_resumo_detalhe_e_retencoes(self):
        from openpyxl import load_workbook
        e1, e2 = suporte.empresa('100'), suporte.empresa('101', cnpj=suporte.CNPJ_B, municipio='3550308')
        r1 = robo.ResultadoEmpresa(e1, extraidas=5, tomadas=3, prestadas=2, fora_do_municipio=2, escrituradas=1,
                                   ja_escrituradas=1, guia_emitida=True,
                                   erros=['nota 3: ISS do portal R$ 50.00 ≠ valor × alíquota R$ 20.00'])
        prestada = nota('7', papel='prestada', tomador_nome='Cliente Exemplo 1',
                        retencoes=robo.Retencoes(inss=110.0, irrf=15.0, csrf=46.5, csrf_quais='PIS+COFINS+CSLL'))
        r1.notas = [(nota('1', retido=True), 'escriturada'), (nota('2'), 'já escriturada'),
                    (nota('3', fonte='SP'), 'NÃO escriturada: ISS do portal R$ 50.00 ≠ valor × alíquota R$ 20.00'),
                    (prestada, 'prestada — com retenção'),
                    (nota('8', papel='prestada', retido=True, cancelada=True), 'prestada cancelada')]
        r2 = robo.ResultadoEmpresa(e2, erros=['RuntimeError: sem certificado A1 válido'])
        with tempfile.TemporaryDirectory() as tmp:
            destino = Path(tmp) / '_relatorios' / 'Analista 1 09-2026.xlsx'     # a pasta ainda não existe
            devolvido = robo.gerar_relatorio([r1, r2], COMP, destino)
            self.assertEqual(devolvido, destino)
            self.assertTrue(destino.exists())
            wb = load_workbook(destino)
            self.assertEqual(wb.sheetnames, ['Empresas', 'Notas', 'Retenções'])
            emp = list(wb['Empresas'].iter_rows(values_only=True))
            det = list(wb['Notas'].iter_rows(values_only=True))
            ret = list(wb['Retenções'].iter_rows(values_only=True))
            self.assertEqual(list(emp[0]), CABECALHO_EMPRESAS)
            self.assertEqual(list(det[0]), CABECALHO_NOTAS)
            self.assertEqual(list(ret[0]), CABECALHO_RETENCOES)
            self.assertEqual(len(emp), 3)                                      # cabeçalho + 2 empresas
            # Retido nas tomadas = ISS da nota 1 (20); nas prestadas = 171,5 da nota 7 (a 8 está cancelada)
            self.assertEqual(list(emp[1]), ['100', e1.razao, e1.cnpj, 5, 3, 2, 2, 1, 1, 'sim', 20.0, 171.5,
                                            'nota 3: ISS do portal R$ 50.00 ≠ valor × alíquota R$ 20.00'])
            self.assertEqual(list(emp[2])[:3], ['101', e2.razao, e2.cnpj])
            self.assertIn(emp[2][9], ('', None))                               # sem guia (célula vazia volta None)
            self.assertEqual(emp[2][12], 'RuntimeError: sem certificado A1 válido')
            self.assertEqual(len(det), 6)                                      # cabeçalho + 5 notas
            linha1 = list(det[1])
            self.assertEqual(linha1[:7], ['100', 'tomada', 'Nacional', '1', '10/09/2026', 'Prestador Exemplo 1',
                                          suporte.CNPJ_PRESTADOR])
            self.assertIn(linha1[7], ('', None))                               # tomador sem nome na nota sintética
            self.assertEqual(linha1[8:14], [suporte.CNPJ_A, suporte.OUTRO_MUNICIPIO, 1000.0, 2.0, 20.0, 'sim'])
            self.assertEqual(linha1[14:20], [0, 0, 0, 0, 0, 0])                # nenhuma retenção federal
            self.assertEqual(linha1[20:], [20.0, 980.0, 'ISS', 'escriturada'])
            self.assertEqual(det[3][2], 'SP')
            self.assertTrue(det[3][23].startswith('NÃO escriturada'))
            linha7 = list(det[4])
            self.assertEqual(linha7[1], 'prestada')
            self.assertEqual(linha7[5:9], ['Empresa Exemplo 100 Ltda', suporte.CNPJ_A, 'Cliente Exemplo 1', suporte.CNPJ_B])
            self.assertEqual(linha7[13:22], ['não', 110.0, 15.0, 0, 0, 0, 46.5, 171.5, 828.5])
            self.assertEqual(linha7[22], 'INSS, IRRF, CSRF (PIS+COFINS+CSLL)')
            # Retenções: uma linha por empresa × papel que tem retenção (a prestada cancelada não conta)
            self.assertEqual(len(ret), 3)
            self.assertEqual(list(ret[1]), ['100', e1.razao, 'tomada', 1, 20.0, 0, 0, 0, 0, 0, 0, 20.0])
            self.assertEqual(list(ret[2]), ['100', e1.razao, 'prestada', 1, 0, 110.0, 15.0, 0, 0, 0, 46.5, 171.5])
            self.assertTrue(wb['Empresas']['A1'].font.bold)                    # cabeçalho formatado
            self.assertGreaterEqual(wb['Notas'].column_dimensions['L'].width, 10)
            wb.close()

    def test_sem_resultados_gera_so_cabecalhos(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as tmp:
            arq = robo.gerar_relatorio([], COMP, Path(tmp) / 'vazio.xlsx')
            wb = load_workbook(arq)
            self.assertEqual(wb['Empresas'].max_row, 1)
            self.assertEqual(wb['Notas'].max_row, 1)
            self.assertEqual(wb['Retenções'].max_row, 1)
            wb.close()


class TestEspelhoPdf(unittest.TestCase):
    def test_segue_sem_reportlab_e_gera_com_ele(self):
        with tempfile.TemporaryDirectory() as tmp:
            destino = Path(tmp) / 'espelho.pdf'
            ok = robo.gerar_espelho_pdf(nota('1'), destino)
            tem_reportlab = importlib.util.find_spec('reportlab') is not None
            self.assertEqual(ok, tem_reportlab)
            self.assertEqual(destino.exists(), tem_reportlab)
            if tem_reportlab:
                self.assertTrue(destino.read_bytes().startswith(b'%PDF'))
                # com retenções e divergência também sai
                r = robo.Retencoes(inss=110.0, csrf=46.5, csrf_quais='PIS+COFINS+CSLL', total_informado=10.0)
                self.assertTrue(robo.gerar_espelho_pdf(nota('2', papel='prestada', retencoes=r), Path(tmp) / 'p.pdf'))

    def test_sem_reportlab_mesmo_que_ele_exista(self):
        """Força o ImportError para garantir o caminho 'espelho é conveniência'."""
        with mock.patch.dict('sys.modules', {'reportlab': None, 'reportlab.lib.pagesizes': None,
                                             'reportlab.pdfgen': None}):
            with tempfile.TemporaryDirectory() as tmp:
                destino = Path(tmp) / 'espelho.pdf'
                self.assertFalse(robo.gerar_espelho_pdf(nota('1'), destino))
                self.assertFalse(destino.exists())


class TestEnviarPorEmail(unittest.TestCase):
    def test_sem_configuracao_nao_envia(self):
        with tempfile.TemporaryDirectory() as tmp, \
                mock.patch.dict(os.environ, {'NFSE_SMTP_SENHA': ''}):
            anexo = Path(tmp) / 'r.xlsx'
            anexo.write_bytes(b'x')
            cfg = robo.Config(pasta_saida=Path(tmp), pasta_estado=Path(tmp), smtp_host='', smtp_usuario='')
            with self.assertLogs('nfse', 'INFO') as logs:
                self.assertFalse(robo.enviar_por_email(cfg, 'alguem@exemplo.test', anexo, COMP))
            self.assertTrue(any('não configurado' in m for m in logs.output))
            # servidor e usuário preenchidos, mas sem a senha no ambiente: também não tenta
            cfg = robo.Config(pasta_saida=Path(tmp), pasta_estado=Path(tmp),
                              smtp_host='smtp.exemplo.test', smtp_usuario='robo@exemplo.test')
            self.assertFalse(robo.enviar_por_email(cfg, 'alguem@exemplo.test', anexo, COMP))
            # sem destinatário, idem
            with mock.patch.dict(os.environ, {'NFSE_SMTP_SENHA': 'x'}):
                self.assertFalse(robo.enviar_por_email(cfg, '', anexo, COMP))


class TestProgresso(unittest.TestCase):
    def test_conta_e_registra(self):
        p = robo.Progresso(2)
        with self.assertLogs('nfse', 'INFO') as logs:
            p.avanca('primeira')
            p.avanca('segunda')
        self.assertEqual(p.feitos, 2)
        self.assertIn('[1/2] primeira', logs.output[0])
        self.assertIn('[2/2] segunda', logs.output[1])
        self.assertIn('faltam ~0s', logs.output[1])


class TestPapeis(unittest.TestCase):
    def test_opcao_so(self):
        self.assertEqual(robo.papeis_de(None), ('tomada', 'prestada'))
        self.assertEqual(robo.papeis_de('tomadas'), ('tomada',))
        self.assertEqual(robo.papeis_de('prestadas'), ('prestada',))


if __name__ == '__main__':
    unittest.main()
