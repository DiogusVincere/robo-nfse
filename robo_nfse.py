# -*- coding: utf-8 -*-
"""
Robô NFS-e — tomadas e prestadas: extração, organização, retenções e escrituração automática
===========================================================================================

Versão de PORTFÓLIO, em arquivo único: mostra a ESTRUTURA e a lógica do robô usado em produção num escritório
contábil. Não contém credenciais, chaves, nomes de clientes, caminhos internos nem endereços de sistemas privados.
Os pontos que dependem de configuração do ambiente estão marcados como CONFIGURÁVEL; o portal municipal (que no robô
real é automatizado por navegador e por requisições HTTP ao JSF da prefeitura) aparece aqui como uma INTERFACE, com
uma implementação simulada para demonstração.

O PROBLEMA
----------
Toda empresa que CONTRATA um serviço (o TOMADOR) precisa conhecer as notas de serviço (NFS-e) emitidas contra ela.
Quando o prestador é de OUTRO município, a prefeitura do tomador não enxerga essa nota — e o tomador é obrigado a
declará-la ("escriturar") no portal da própria prefeitura, recolhendo o ISS quando ele é retido. Feito à mão, isso
significa abrir portal por portal, empresa por empresa, todo mês.

A mesma empresa, quando PRESTA serviço, precisa saber o que cada cliente RETEVE do pagamento — ISS, INSS, IRRF,
PIS/COFINS/CSLL — porque é esse valor que ela abate na apuração do mês. A informação está nota a nota, em campos
cujo significado mudou em 2026 (ver seção 6), e somar errado vira imposto pago em dobro ou a menos.

O QUE O ROBÔ FAZ (uma execução, ponta a ponta)
----------------------------------------------
  ETAPA 1  EXTRAÇÃO, 100% em segundo plano (sem navegador), com o certificado A1 de cada empresa, das notas que ela
           TOMOU e das que ela PRESTOU:
           • Prefeitura de São Paulo (Nota Paulistana) — web service SOAP, pedido assinado (XMLDSig) + mTLS:
             ConsultaNFeRecebidas (tomadas) e ConsultaNFeEmitidas (prestadas);
           • Sistema Nacional NFS-e (ADN) — API REST de distribuição de documentos por NSU, mTLS, incremental
             (guarda o último NSU por CNPJ), com controle de rate limit e backoff. A distribuição devolve tudo em
             que o CNPJ aparece; o robô separa pelo papel (tomador = tomada, emitente = prestada).
  ETAPA 2  ORGANIZAÇÃO: cada XML vai para a pasta do cliente na competência certa (competência = mês da EMISSÃO),
           em Tomados/ ou Prestados/, com espelho em PDF; documentos de outras competências são arquivados.
  ETAPA 3  CLASSIFICAÇÃO E RETENÇÕES: as tomadas pelo município do PRESTADOR (código IBGE) — "do município" ×
           "fora do município", autorizadas × canceladas, deduplicadas entre as duas fontes; e, em toda nota, o
           que foi RETIDO: ISS (e por quem), INSS, IRRF, PIS, COFINS, CSLL — ou a soma PIS+COFINS+CSLL quando o
           leiaute só traz a soma (NT SE/CGNFS-e 007; Nota Paulistana desde 14/05/2026).
  ETAPA 4  ESCRITURAÇÃO no portal municipal das tomadas "fora do município":
           • consulta ANTI-DUPLICIDADE antes (o que já foi escriturado não é emitido de novo → re-executável);
           • TRAVAS antes de cada emissão: número do documento confere; o ISS recalculado (valor × alíquota) bate;
           • salva o comprovante (PDF) e consulta de novo para marcar "escriturada";
           • sessão monitorada: se o portal derrubar, renova e continua de onde parou.
  ETAPA 5  PÓS-ESCRITURAÇÃO: guia de ISS retido, exportação dos XML para o sistema contábil (com a renumeração que
           o portal exige) e relatório por responsável — com uma aba de retenções por tributo —, enviado por e-mail.

DECISÕES DE PROJETO QUE VALEM A PENA DESTACAR
---------------------------------------------
  • Idempotência: toda etapa pode ser repetida sem duplicar nada (estado de NSU, consulta antes de emitir, ledger).
  • Falha isolada: erro numa empresa vira linha no relatório e o robô segue para a próxima.
  • Segurança: a chave privada do A1 só existe em PEM temporário, apagado ao sair do contexto; senhas vêm de um
    provedor configurável (variável de ambiente, cofre etc.), nunca do código.
  • Gentileza com os serviços públicos: pausas entre chamadas, backoff em 429/5xx, respeito ao Retry-After.
  • Retenção é LIDA, nunca deduzida: o robô mostra o que o XML diz e avisa quando o total informado não fecha.

COMO RODAR
----------
    pip install requests lxml cryptography openpyxl
    python robo_nfse.py --demo                          # fluxo completo com dados FICTÍCIOS, offline
    python robo_nfse.py --interface                     # página local no navegador: certificado, competência, zip
    python robo_nfse.py --competencia 09/2026 --empresas empresas.json --certificados ./certs --saida ./saida
    python robo_nfse.py ... --so prestadas              # só um dos papéis (padrão: tomadas e prestadas)

Estrutura do código (no robô de produção cada seção é um módulo; aqui estão em ordem):
  1 Configuração e utilitários      6 Leitura dos XML → Nota (+ retenções)  11 Exportação para o sistema contábil
  2 Modelos                         7 Pastas e estado (NSU)                 12 Relatório e e-mail
  3 Certificado A1 e mTLS           8 Espelho em PDF (opcional)             13 Orquestrador
  4 Assinatura XML (XMLDSig)        9 Classificação                         14 Demonstração e linha de comando
  5 Clientes dos web services      10 Escrituração (interface + travas)
"""
from __future__ import annotations

import abc
import argparse
import base64
import contextlib
import dataclasses
import datetime as dt
import gzip
import hashlib
import json
import logging
import os
import random
import re
import shutil
import smtplib
import tempfile
import time
from email.message import EmailMessage
from pathlib import Path
from typing import Callable, Iterable, Iterator, Optional

from lxml import etree

log = logging.getLogger('nfse')

PAPEIS = ('tomada', 'prestada')                         # os dois lados de uma NFS-e, do ponto de vista da empresa
PASTA_DO_PAPEL = {'tomada': 'Tomados', 'prestada': 'Prestados'}


# =====================================================================================================================
# 1. CONFIGURAÇÃO E UTILITÁRIOS
# =====================================================================================================================
@dataclasses.dataclass
class Config:
    """Tudo o que muda de um ambiente para outro. Nada aqui é segredo: senhas e chaves vêm de provedores (seção 3)."""
    pasta_saida: Path                                   # raiz das pastas dos clientes (CONFIGURÁVEL)
    pasta_estado: Path                                  # estado do robô: último NSU, arquivo, ledger
    municipio_escrituracao_ibge: str = '3509502'        # código IBGE do município do portal de escrituração (CONFIGURÁVEL)
    pausa_entre_empresas: float = 2.0                   # gentileza com os web services
    tolerancia_iss: float = 0.01                        # R$ — diferença aceita entre ISS do portal e o calculado
    papeis: tuple = PAPEIS                              # o que buscar: notas que a empresa TOMOU e/ou que ela PRESTOU
    smtp_host: str = os.environ.get('NFSE_SMTP_HOST', '')
    smtp_porta: int = int(os.environ.get('NFSE_SMTP_PORTA', '587') or 587)
    smtp_usuario: str = os.environ.get('NFSE_SMTP_USUARIO', '')
    # a senha do SMTP NÃO fica na configuração: é lida do ambiente na hora do envio (NFSE_SMTP_SENHA)


def digitos(v) -> str:
    return re.sub(r'\D', '', str(v or ''))


def local(el, nome: str) -> Optional[str]:
    """Texto do primeiro descendente com este NOME LOCAL, ignorando namespace — os layouts municipais variam muito."""
    if el is None:
        return None
    achado = next((e for e in el.iter() if isinstance(e.tag, str) and etree.QName(e).localname == nome), None)
    return (achado.text or '').strip() if achado is not None and achado.text else None


def grupo(el, *nomes: str):
    """O primeiro elemento cujo nome local está em `nomes` (ex.: 'emit', 'prest', 'prestador')."""
    if el is None:
        return None
    return next((e for e in el.iter() if isinstance(e.tag, str) and etree.QName(e).localname in nomes), None)


def para_float(v) -> float:
    if v in (None, ''):
        return 0.0
    s = str(v).strip()
    if ',' in s and '.' in s:
        s = s.replace('.', '').replace(',', '.')
    elif ',' in s:
        s = s.replace(',', '.')
    try:
        return float(s)
    except ValueError:
        return 0.0


@dataclasses.dataclass(frozen=True, order=True)
class Competencia:
    ano: int
    mes: int

    @classmethod
    def de_texto(cls, txt: str) -> 'Competencia':
        """'09/2026', '09-2026' ou '2026-09'."""
        m = re.fullmatch(r'(\d{1,2})[/-](\d{4})', txt.strip()) or None
        if m:
            return cls(int(m.group(2)), int(m.group(1)))
        m = re.fullmatch(r'(\d{4})-(\d{1,2})', txt.strip())
        if m:
            return cls(int(m.group(1)), int(m.group(2)))
        raise ValueError(f'competência inválida: {txt!r} (use MM/AAAA)')

    @classmethod
    def anterior_a(cls, hoje: dt.date) -> 'Competencia':
        d = hoje.replace(day=1) - dt.timedelta(days=1)
        return cls(d.year, d.month)

    @property
    def inicio(self) -> dt.date:
        return dt.date(self.ano, self.mes, 1)

    @property
    def fim(self) -> dt.date:
        prox = (self.inicio.replace(day=28) + dt.timedelta(days=4)).replace(day=1)
        return prox - dt.timedelta(days=1)

    def __str__(self):
        return f'{self.mes:02d}/{self.ano}'

    @property
    def pasta(self) -> str:
        return f'{self.mes:02d}-{self.ano}'


def nome_seguro(txt: str, limite: int = 80) -> str:
    """Nome de arquivo/pasta válido no Windows e no Linux."""
    txt = re.sub(r'[<>:"/\\|?*\x00-\x1f]', ' ', str(txt or '')).strip().rstrip('.')
    return re.sub(r'\s+', ' ', txt)[:limite] or 'sem_nome'


# =====================================================================================================================
# 2. MODELOS
# =====================================================================================================================
@dataclasses.dataclass
class Empresa:
    """Um cliente do escritório — tomador das notas que recebe e prestador das que emite.
    A lista vem de um cadastro central (CONFIGURÁVEL: JSON, API, planilha)."""
    codigo: str
    razao: str
    cnpj: str
    municipio_ibge: str
    responsavel: str                                    # quem cuida da empresa no departamento fiscal
    pasta: Optional[Path] = None                        # pasta do cliente (resolvida pela seção 7)


# tpRetPisCofins do leiaute nacional → quais contribuições foram retidas. Os códigos 1 e 2 são da regra ANTIGA
# (os valores retidos de PIS/COFINS vinham em vPis/vCofins); de 3 a 9 valem com a NT SE/CGNFS-e 007, em que vRetCSLL
# passa a ser a SOMA das retenções de PIS, COFINS e CSLL e o código diz quais entraram na soma.
CODIGOS_RET_PIS_COFINS = {
    '0': (), '2': (),
    '1': ('PIS', 'COFINS'),
    '3': ('PIS', 'COFINS', 'CSLL'), '4': ('PIS', 'COFINS'), '5': ('PIS',), '6': ('COFINS',),
    '7': ('COFINS', 'CSLL'), '8': ('CSLL',), '9': ('PIS', 'CSLL'),
}
# Nota Paulistana adaptada ao modelo nacional: a partir desta data ValorPIS/ValorCOFINS são o débito próprio do
# prestador e ValorCSLL é a SOMA das retenções de PIS, COFINS e CSLL (antes, cada campo era a própria retenção).
SP_CSRF_SOMADA_DESDE = dt.date(2026, 5, 14)


@dataclasses.dataclass
class Retencoes:
    """O que o TOMADOR descontou do pagamento para recolher em nome do prestador, em R$, como está no XML.

    ISS: tpRetISSQN 2 (tomador) ou 3 (intermediário) no Nacional; ISSRetido=true em SP.
    INSS (contribuição previdenciária): vRetCP / ValorINSS.  IRRF: vRetIRRF / ValorIR.
    PIS, COFINS e CSLL: separados quando o leiaute separa; quando ele só traz a soma (NT 007, SP desde 14/05/2026),
    a soma fica em `csrf` e `csrf_quais` diz o que está dentro dela."""
    iss: float = 0.0
    iss_por: str = ''                                   # 'tomador' ou 'intermediário'
    inss: float = 0.0
    irrf: float = 0.0
    pis: float = 0.0
    cofins: float = 0.0
    csll: float = 0.0
    csrf: float = 0.0                                   # PIS+COFINS+CSLL somados, quando o XML não separa
    csrf_quais: str = ''                                # 'PIS+COFINS+CSLL', 'PIS+COFINS', 'COFINS+CSLL', 'PIS+CSLL'
    total_informado: Optional[float] = None             # vTotalRet do Nacional, para conferência

    @property
    def federais(self) -> float:
        return round(self.inss + self.irrf + self.pis + self.cofins + self.csll + self.csrf, 2)

    @property
    def total(self) -> float:
        return round(self.iss + self.federais, 2)

    @property
    def tem(self) -> bool:
        return self.total > 0

    @property
    def do_tomador(self) -> float:
        """O que o próprio TOMADOR reteve e vai recolher: as federais e o ISS — menos o ISS retido pelo intermediário,
        que é ele quem recolhe. É a parte de `total` que entra em 'a empresa recolhe' nas tomadas."""
        return round(self.federais + (0.0 if self.iss_por == 'intermediário' else self.iss), 2)

    def quais(self) -> list[str]:
        """Nomes do que foi retido, na ordem em que o fiscal confere: ['ISS', 'INSS', 'IRRF', 'CSRF (PIS+COFINS+CSLL)']."""
        pares = [('ISS', self.iss), ('INSS', self.inss), ('IRRF', self.irrf), ('PIS', self.pis),
                 ('COFINS', self.cofins), ('CSLL', self.csll),
                 (f'CSRF ({self.csrf_quais})' if self.csrf_quais else 'CSRF', self.csrf)]
        return [nome for nome, valor in pares if valor > 0]

    def divergencia(self) -> Optional[str]:
        """O Nacional informa o total das retenções (vTotalRet): se não bater com a soma lida, alguém errou — avisa."""
        if self.total_informado is None or abs(self.total_informado - self.total) <= 0.01:
            return None
        return f'vTotalRet R$ {self.total_informado:.2f} ≠ soma das retenções R$ {self.total:.2f}'

    def descricao(self) -> str:
        texto = ', '.join(self.quais()) or 'nenhuma'
        aviso = self.divergencia()
        return f'{texto} — conferir: {aviso}' if aviso else texto


@dataclasses.dataclass
class Nota:
    """Uma NFS-e já normalizada — o mesmo formato para as duas fontes e para os dois papéis."""
    fonte: str                                          # 'SP' ou 'Nacional'
    numero: str
    chave: str                                          # chave de acesso (Nacional) ou nº + código de verificação (SP)
    emissao: Optional[dt.date]
    competencia: Optional[Competencia]                  # = mês da EMISSÃO (regra do fiscal)
    prestador_cnpj: str
    prestador_nome: str
    prestador_municipio_ibge: str
    tomador_cnpj: str
    valor_servico: float
    aliquota: float                                     # em %, ex.: 2.0
    valor_iss: float
    # `iss_retido` e `valor_liquido` entram no construtor por comodidade (demonstração, testes, vLiq do XML), mas
    # NÃO são guardados: `retencoes` é a única fonte da verdade — ver as propriedades logo abaixo da classe.
    iss_retido: dataclasses.InitVar[bool] = False
    cancelada: bool = False
    xml: str = dataclasses.field(default='', repr=False)
    papel: str = 'tomada'                               # 'tomada' (a empresa é o tomador) ou 'prestada' (é o prestador)
    tomador_nome: str = ''
    tomador_municipio_ibge: str = ''
    retencoes: Retencoes = dataclasses.field(default_factory=Retencoes)
    valor_liquido: dataclasses.InitVar[Optional[float]] = None
    valor_liquido_lido: Optional[float] = dataclasses.field(default=None, repr=False)   # vLiq do Nacional, se veio

    def __post_init__(self, iss_retido: bool = False, valor_liquido: Optional[float] = None):
        r = self.retencoes
        if r.iss or r.iss_por:                           # quem leu o XML preencheu a retenção: ela é a fonte da verdade
            r.iss_por = r.iss_por or 'tomador'
        elif iss_retido:                                 # quem monta a nota só com o booleano (demonstração, testes)
            r.iss, r.iss_por = round(self.valor_iss, 2), 'tomador'
        if valor_liquido is not None:
            self.valor_liquido_lido = valor_liquido

    @property
    def identidade(self) -> tuple:
        """Para deduplicar a mesma nota vinda das duas fontes."""
        return (self.prestador_cnpj, self.numero.lstrip('0'), round(self.valor_servico, 2))

    @property
    def contraparte(self) -> str:
        """Quem está do outro lado: o prestador numa tomada, o tomador numa prestada (dá nome ao arquivo)."""
        return self.tomador_nome if self.papel == 'prestada' else self.prestador_nome


def _iss_retido(self: Nota) -> bool:
    return bool(self.retencoes.iss or self.retencoes.iss_por)


def _definir_iss_retido(self: Nota, valor: bool) -> None:
    if valor:
        self.retencoes.iss, self.retencoes.iss_por = round(self.valor_iss, 2), self.retencoes.iss_por or 'tomador'
    else:
        self.retencoes.iss, self.retencoes.iss_por = 0.0, ''


def _valor_liquido(self: Nota) -> float:
    """O vLiq do XML quando veio; senão, valor do serviço − tudo o que foi retido. Sempre calculado a partir de
    `retencoes`, para não desalinhar quando alguém troca as retenções depois de montar a nota."""
    if self.valor_liquido_lido is not None:
        return self.valor_liquido_lido
    return round(self.valor_servico - self.retencoes.total, 2)


def _definir_valor_liquido(self: Nota, valor: Optional[float]) -> None:
    self.valor_liquido_lido = valor


# As propriedades são ligadas DEPOIS do decorador: dentro da classe, o nome `iss_retido` é o InitVar do construtor
# (um atributo de classe com o mesmo nome viraria o "default" do parâmetro).
Nota.iss_retido = property(_iss_retido, _definir_iss_retido)
Nota.valor_liquido = property(_valor_liquido, _definir_valor_liquido)


@dataclasses.dataclass
class ResultadoEmpresa:
    empresa: Empresa
    extraidas: int = 0                                  # documentos lidos (antes de deduplicar)
    tomadas: int = 0
    prestadas: int = 0
    fora_do_municipio: int = 0
    escrituradas: int = 0
    ja_escrituradas: int = 0
    guia_emitida: bool = False
    erros: list = dataclasses.field(default_factory=list)
    notas: list = dataclasses.field(default_factory=list)   # [(Nota, situação)]

    def retido(self, papel: str) -> float:
        """Soma das retenções das notas ativas de um papel — o que a empresa recolhe (tomadas: só o que ela mesma
        reteve, sem o ISS retido pelo intermediário) ou abate (prestadas: tudo o que os clientes retiveram)."""
        return round(sum(retido_pela_empresa(n) for n, _ in self.notas if n.papel == papel and not n.cancelada), 2)


def retido_pela_empresa(n: Nota) -> float:
    """Numa tomada, o que a empresa (tomadora) reteve e recolhe; numa prestada, tudo o que retiveram dela."""
    return n.retencoes.do_tomador if n.papel != 'prestada' else n.retencoes.total


def iss_retido_pela_empresa(n: Nota) -> float:
    """ISS que a empresa recolhe por ter retido numa tomada (o retido pelo intermediário não é dela)."""
    return n.retencoes.iss if n.papel != 'prestada' and n.retencoes.iss_por != 'intermediário' else 0.0


# =====================================================================================================================
# 3. CERTIFICADO A1 E mTLS
# =====================================================================================================================
class SenhaProvider(abc.ABC):
    """De onde vem a senha do .pfx. No robô real: um cofre central. NUNCA no código nem no nome do arquivo."""

    @abc.abstractmethod
    def senha(self, cnpj: str) -> str:
        ...


class SenhaDoAmbiente(SenhaProvider):
    """Exemplo simples: variável de ambiente NFSE_PFX_SENHA_<CNPJ>."""

    def senha(self, cnpj: str) -> str:
        valor = os.environ.get(f'NFSE_PFX_SENHA_{digitos(cnpj)}')
        if not valor:
            raise LookupError(f'senha do certificado de {cnpj} não configurada')
        return valor


@dataclasses.dataclass
class CertificadoA1:
    caminho: Path
    cnpj: str
    titular: str
    validade: dt.datetime


OID_CNPJ_ICP_BRASIL = '2.16.76.1.3.3'                    # otherName do e-CNPJ (ICP-Brasil)


def ler_certificado(caminho: Path, senha: str) -> CertificadoA1:
    """Abre o .pfx e lê o CNPJ DE DENTRO do certificado — o nome da pasta/arquivo não é confiável."""
    from cryptography import x509
    from cryptography.hazmat.primitives.serialization import pkcs12
    _chave, cert, _cadeia = pkcs12.load_key_and_certificates(caminho.read_bytes(), senha.encode() if senha else None)
    if cert is None:
        raise ValueError('pfx sem certificado')
    cn = next((a.value for a in cert.subject if a.oid == x509.NameOID.COMMON_NAME), '')
    cnpj = ''
    m = re.search(r':(\d{14})\b', cn)                    # padrão ICP-Brasil: "RAZÃO SOCIAL:12345678000195"
    if m:
        cnpj = m.group(1)
    else:
        with contextlib.suppress(x509.ExtensionNotFound):
            san = cert.extensions.get_extension_for_class(x509.SubjectAlternativeName).value
            for outro in san.get_values_for_type(x509.OtherName):
                if outro.type_id.dotted_string == OID_CNPJ_ICP_BRASIL:
                    achado = re.search(rb'\d{14}', outro.value)
                    cnpj = achado.group().decode() if achado else ''
    if len(cnpj) != 14:
        raise ValueError(f'não achei o CNPJ dentro de {caminho.name}')
    validade = getattr(cert, 'not_valid_after_utc', None) or cert.not_valid_after.replace(tzinfo=dt.timezone.utc)
    return CertificadoA1(caminho, cnpj, cn.split(':')[0], validade)


def certificados_por_cnpj(pasta: Path, senhas: SenhaProvider, cnpjs: Iterable[str]) -> dict[str, CertificadoA1]:
    """Varre a pasta de certificados (uma subpasta por cliente) e devolve o .pfx MAIS RECENTE e VÁLIDO de cada CNPJ."""
    desejados = {digitos(c) for c in cnpjs}
    achados: dict[str, CertificadoA1] = {}
    for pfx in sorted(pasta.rglob('*.pfx'), key=lambda p: p.stat().st_mtime, reverse=True):
        for cnpj in desejados - set(achados):
            try:
                cert = ler_certificado(pfx, senhas.senha(cnpj))
            except Exception:                                          # senha de outro CNPJ: não é este
                continue
            if cert.cnpj == cnpj and cert.validade > dt.datetime.now(dt.timezone.utc):
                achados[cnpj] = cert
                break
    return achados


@contextlib.contextmanager
def sessao_mtls(cert: CertificadoA1, senha: str) -> Iterator['requests.Session']:
    """Sessão `requests` com autenticação mútua. O PEM (com a chave privada) existe só durante o `with`."""
    import requests
    from cryptography.hazmat.primitives.serialization import Encoding, NoEncryption, PrivateFormat, pkcs12
    chave, c, cadeia = pkcs12.load_key_and_certificates(cert.caminho.read_bytes(), senha.encode() if senha else None)
    cert_pem = c.public_bytes(Encoding.PEM) + b''.join(ca.public_bytes(Encoding.PEM) for ca in (cadeia or []))
    key_pem = chave.private_bytes(Encoding.PEM, PrivateFormat.TraditionalOpenSSL, NoEncryption())
    pasta = Path(tempfile.mkdtemp(prefix='mtls_'))
    (pasta / 'cert.pem').write_bytes(cert_pem)
    (pasta / 'key.pem').write_bytes(key_pem)
    sess = requests.Session()
    sess.cert = (str(pasta / 'cert.pem'), str(pasta / 'key.pem'))
    try:
        yield sess
    finally:
        sess.close()
        shutil.rmtree(pasta, ignore_errors=True)                      # a chave privada nunca fica no disco


# =====================================================================================================================
# 4. ASSINATURA XML (XMLDSig enveloped, RSA-SHA1) — exigida pelo web service da Prefeitura de São Paulo
# =====================================================================================================================
DS = 'http://www.w3.org/2000/09/xmldsig#'


def assinar_enveloped(xml: str, cert: CertificadoA1, senha: str) -> str:
    """Reference URI="" + transforms enveloped-signature e C14N; KeyInfo só com o certificado do titular."""
    from cryptography.hazmat.primitives import hashes
    from cryptography.hazmat.primitives.asymmetric import padding
    from cryptography.hazmat.primitives.serialization import Encoding, pkcs12
    chave, c, _ = pkcs12.load_key_and_certificates(cert.caminho.read_bytes(), senha.encode() if senha else None)
    c14n = lambda el: etree.tostring(el, method='c14n', exclusive=False, with_comments=False)   # noqa: E731
    ds = lambda tag: '{%s}%s' % (DS, tag)                                                         # noqa: E731
    raiz = etree.fromstring(xml.encode('utf-8'))
    digest = base64.b64encode(hashlib.sha1(c14n(raiz)).digest()).decode()
    sig = etree.SubElement(raiz, ds('Signature'))
    info = etree.SubElement(sig, ds('SignedInfo'))
    etree.SubElement(info, ds('CanonicalizationMethod')).set('Algorithm', 'http://www.w3.org/TR/2001/REC-xml-c14n-20010315')
    etree.SubElement(info, ds('SignatureMethod')).set('Algorithm', DS + 'rsa-sha1')
    ref = etree.SubElement(info, ds('Reference'))
    ref.set('URI', '')
    trs = etree.SubElement(ref, ds('Transforms'))
    etree.SubElement(trs, ds('Transform')).set('Algorithm', DS + 'enveloped-signature')
    etree.SubElement(trs, ds('Transform')).set('Algorithm', 'http://www.w3.org/TR/2001/REC-xml-c14n-20010315')
    etree.SubElement(ref, ds('DigestMethod')).set('Algorithm', DS + 'sha1')
    etree.SubElement(ref, ds('DigestValue')).text = digest
    assinatura = chave.sign(c14n(info), padding.PKCS1v15(), hashes.SHA1())
    etree.SubElement(sig, ds('SignatureValue')).text = base64.b64encode(assinatura).decode()
    x509d = etree.SubElement(etree.SubElement(sig, ds('KeyInfo')), ds('X509Data'))
    etree.SubElement(x509d, ds('X509Certificate')).text = base64.b64encode(c.public_bytes(Encoding.DER)).decode()
    return etree.tostring(raiz, xml_declaration=True, encoding='UTF-8').decode('utf-8')


# =====================================================================================================================
# 5. CLIENTES DOS WEB SERVICES (públicos e documentados)
# =====================================================================================================================
class ClienteSaoPaulo:
    """Nota Paulistana — ConsultaNFeRecebidas (notas que a empresa TOMOU) e ConsultaNFeEmitidas (que ela PRESTOU),
    por período. SOAP 1.1, pedido assinado, mTLS com o A1 da empresa. A mensagem é a mesma nas duas operações
    (PedidoConsultaNFePeriodo); muda o nome da operação e o SOAPAction.
    Paginação: até 50 NF-e por página; segue até uma página vir com menos de 50."""
    NS = 'http://www.prefeitura.sp.gov.br/nfe'
    ENDPOINT = 'https://nfe.prefeitura.sp.gov.br/ws/lotenfe.asmx'
    SOAP = 'http://schemas.xmlsoap.org/soap/envelope/'
    POR_PAGINA = 50
    OPERACOES = {'tomada': 'ConsultaNFeRecebidas', 'prestada': 'ConsultaNFeEmitidas'}

    def _pedido(self, cnpj: str, ini: dt.date, fim: dt.date, pagina: int) -> str:
        # schema "unqualified": só a raiz no namespace; os filhos sem namespace
        raiz = etree.Element('{%s}PedidoConsultaNFePeriodo' % self.NS, nsmap={'nfe': self.NS})
        cab = etree.SubElement(raiz, 'Cabecalho')
        cab.set('Versao', '1')
        etree.SubElement(etree.SubElement(cab, 'CPFCNPJRemetente'), 'CNPJ').text = cnpj
        etree.SubElement(etree.SubElement(cab, 'CPFCNPJ'), 'CNPJ').text = cnpj
        etree.SubElement(cab, 'dtInicio').text = ini.isoformat()
        etree.SubElement(cab, 'dtFim').text = fim.isoformat()
        etree.SubElement(cab, 'NumeroPagina').text = str(pagina)
        return etree.tostring(raiz, xml_declaration=True, encoding='UTF-8').decode('utf-8')

    def _envelope(self, mensagem: str, operacao: str = 'ConsultaNFeRecebidas') -> bytes:
        env = etree.Element('{%s}Envelope' % self.SOAP, nsmap={'soap': self.SOAP})
        req = etree.SubElement(etree.SubElement(env, '{%s}Body' % self.SOAP),
                               '{%s}%sRequest' % (self.NS, operacao), nsmap={None: self.NS})
        etree.SubElement(req, '{%s}VersaoSchema' % self.NS).text = '1'
        etree.SubElement(req, '{%s}MensagemXML' % self.NS).text = mensagem         # vai como TEXTO (escapado)
        return etree.tostring(env, xml_declaration=True, encoding='UTF-8')

    @classmethod
    def soapaction(cls, operacao: str) -> str:
        """'ConsultaNFeRecebidas' → http://www.prefeitura.sp.gov.br/nfe/ws/consultaNFeRecebidas (inicial minúscula)."""
        return f'{cls.NS}/ws/{operacao[0].lower()}{operacao[1:]}'

    def consultar(self, cert: CertificadoA1, senha: str, comp: Competencia, papel: str = 'tomada') -> list[str]:
        operacao = self.OPERACOES[papel]
        xmls: list[str] = []
        with sessao_mtls(cert, senha) as sess:
            for pagina in range(1, 201):
                pedido = assinar_enveloped(self._pedido(cert.cnpj, comp.inicio, comp.fim, pagina), cert, senha)
                r = sess.post(self.ENDPOINT, data=self._envelope(pedido, operacao), timeout=240,
                              headers={'Content-Type': 'text/xml; charset=utf-8', 'SOAPAction': self.soapaction(operacao)})
                r.raise_for_status()
                corpo = etree.fromstring(r.content)
                retorno = grupo(corpo, 'RetornoXML')
                if retorno is None or not (retorno.text or '').strip():
                    raise RuntimeError('resposta SOAP sem RetornoXML')
                retorno = etree.fromstring(retorno.text.encode('utf-8'))
                if (local(retorno, 'Sucesso') or '').lower() == 'false':
                    raise RuntimeError(f"web service de SP: [{local(retorno, 'Codigo')}] {local(retorno, 'Descricao')}")
                nfes = [e for e in retorno.iter() if isinstance(e.tag, str) and etree.QName(e).localname == 'NFe']
                xmls += [etree.tostring(n, encoding='UTF-8').decode('utf-8') for n in nfes]
                log.info('Nota Paulistana (%s): página %d, %d nota(s)', papel + 's', pagina, len(nfes))
                if len(nfes) < self.POR_PAGINA:
                    break
        return xmls

    def consultar_recebidas(self, cert: CertificadoA1, senha: str, comp: Competencia) -> list[str]:
        return self.consultar(cert, senha, comp, 'tomada')

    def consultar_emitidas(self, cert: CertificadoA1, senha: str, comp: Competencia) -> list[str]:
        return self.consultar(cert, senha, comp, 'prestada')


class ClienteNacional:
    """Sistema Nacional NFS-e (ADN) — distribuição de documentos por NSU, mTLS com o A1 da empresa.
    GET {BASE}/DFe/{NSU}?lote=true&cnpjConsulta=... → até 50 documentos (XML em base64, às vezes gzip), TUDO em que
    o CNPJ aparece: como tomador, como emitente/prestador ou como intermediário (quem separa é a seção 13).
    Incremental: pede a partir do último NSU conhecido. Tem rate limit — pausa fixa + backoff em 429/5xx."""
    BASE = 'https://adn.nfse.gov.br/contribuintes'
    LOTE = 50
    PAUSA = 3
    BACKOFF = (15, 30, 60, 120)
    REPETE = {429, 502, 503, 504}

    @staticmethod
    def _decodifica(b64: str) -> str:
        bruto = base64.b64decode(b64)
        if bruto[:2] == b'\x1f\x8b':                                   # gzip
            bruto = gzip.decompress(bruto)
        return bruto.decode('utf-8', errors='replace')

    def distribuir(self, cert: CertificadoA1, senha: str, ultimo_nsu: int) -> tuple[list[dict], int]:
        docs, nsu = [], ultimo_nsu
        log.info('Sistema Nacional: distribuição a partir do NSU %d%s', ultimo_nsu + 1,
                 '' if ultimo_nsu else ' (primeira vez: percorre o histórico inteiro, pode demorar)')
        with sessao_mtls(cert, senha) as sess:
            tentativas = 0
            for chamada in range(400):
                if chamada:
                    time.sleep(self.PAUSA)
                try:
                    r = sess.get(f'{self.BASE}/DFe/{nsu + 1}', params={'lote': 'true', 'cnpjConsulta': cert.cnpj},
                                 timeout=120)
                    status = r.status_code
                except Exception as e:                                   # queda de rede = como um 5xx
                    r, status = None, f'rede: {e}'
                if r is None or status in self.REPETE:
                    if tentativas >= len(self.BACKOFF):
                        raise RuntimeError(f'ADN indisponível ({status}) — o que já veio está salvo; tente depois')
                    espera = self.BACKOFF[tentativas]
                    with contextlib.suppress(Exception):
                        espera = max(espera, int(r.headers.get('Retry-After') or 0))
                    log.warning('Sistema Nacional respondeu %s — aguardando %ds antes de tentar de novo', status, espera)
                    time.sleep(espera)
                    tentativas += 1
                    continue
                tentativas = 0
                if status == 404:                                        # nada novo a partir deste NSU
                    log.info('Sistema Nacional: nada novo a partir do NSU %d', nsu + 1)
                    break
                r.raise_for_status()
                corpo = r.json()
                if corpo.get('Erros'):
                    raise RuntimeError(f"ADN: {corpo['Erros']}")
                lote = corpo.get('LoteDFe') or []
                for item in lote:
                    nsu = max(nsu, int(item.get('NSU') or 0))
                    docs.append({'nsu': int(item.get('NSU') or 0), 'chave': item.get('ChaveAcesso'),
                                 'tipo': str(item.get('TipoDocumento') or ''),
                                 'xml': self._decodifica(item.get('ArquivoXml') or '')})
                log.info('Sistema Nacional: lote até o NSU %d, %d documento(s) até agora', nsu, len(docs))
                if len(lote) < self.LOTE:
                    break
        return docs, nsu


# =====================================================================================================================
# 6. LEITURA DOS XML → Nota (os dois layouts num formato só) — inclusive o que foi RETIDO
# =====================================================================================================================
# Onde cada retenção mora:
#   Nacional (DPS/infDPS/valores/trib): tribMun/tpRetISSQN (1 não retido, 2 pelo tomador, 3 pelo intermediário);
#     tribFed/vRetCP (INSS), vRetIRRF, vRetCSLL e tribFed/piscofins/tpRetPisCofins. Pela NT SE/CGNFS-e 007, vRetCSLL
#     é a SOMA das retenções de PIS, COFINS e CSLL e o código 3..9 diz quais estão nela; na regra antiga (código 1)
#     PIS e COFINS retidos vinham em vPis/vCofins e vRetCSLL era só a CSLL. infNFSe/valores/vTotalRet traz o total
#     (ISS retido + federais) e vLiq o líquido — servem de conferência.
#   SP (Nota Paulistana): ISSRetido, ValorINSS, ValorIR, ValorPIS, ValorCOFINS, ValorCSLL. Até 13/05/2026 cada campo
#     era a própria retenção; desde 14/05/2026 (adaptação ao modelo nacional) ValorPIS/ValorCOFINS passaram a ser o
#     débito próprio e ValorCSLL a soma das retenções de PIS, COFINS e CSLL. O XML lido não diz quais compõem a soma.
def _data(v: Optional[str]) -> Optional[dt.date]:
    m = re.match(r'(\d{4})-(\d{2})-(\d{2})', v or '')
    return dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3))) if m else None


def retencoes_de_sp(raiz, emissao: Optional[dt.date], valor_iss: float, iss_retido: bool) -> Retencoes:
    r = Retencoes(iss=round(valor_iss, 2) if iss_retido else 0.0, iss_por='tomador' if iss_retido else '')
    r.inss = para_float(local(raiz, 'ValorINSS'))
    r.irrf = para_float(local(raiz, 'ValorIR'))
    pis, cofins, csll = (para_float(local(raiz, c)) for c in ('ValorPIS', 'ValorCOFINS', 'ValorCSLL'))
    if emissao and emissao >= SP_CSRF_SOMADA_DESDE:
        if csll:                                       # a soma PIS+COFINS+CSLL; ValorPIS/ValorCOFINS são débito próprio
            r.csrf, r.csrf_quais = csll, 'PIS+COFINS+CSLL'
    else:
        r.pis, r.cofins, r.csll = pis, cofins, csll
    return r


def retencoes_nacional(raiz, valor_iss: float, tp_ret_iss: str) -> Retencoes:
    r = Retencoes()
    if tp_ret_iss in ('2', '3'):
        r.iss, r.iss_por = round(valor_iss, 2), 'tomador' if tp_ret_iss == '2' else 'intermediário'
    fed = grupo(raiz, 'tribFed')
    if fed is not None:
        r.inss = para_float(local(fed, 'vRetCP'))
        r.irrf = para_float(local(fed, 'vRetIRRF'))
        v_csll = para_float(local(fed, 'vRetCSLL'))
        codigo = local(fed, 'tpRetPisCofins') or ''
        quais = CODIGOS_RET_PIS_COFINS.get(codigo, ())
        if codigo == '1':                              # regra antiga: PIS/COFINS retidos em vPis/vCofins; vRetCSLL só CSLL
            r.pis, r.cofins, r.csll = para_float(local(fed, 'vPis')), para_float(local(fed, 'vCofins')), v_csll
        elif len(quais) == 1 and v_csll:               # 5, 6, 8: a "soma" é de um tributo só — vai para o campo dele
            setattr(r, quais[0].lower(), v_csll)
        elif len(quais) > 1 and v_csll:                # 3, 4, 7, 9: NT 007 — vRetCSLL é a soma; o código diz de quais
            r.csrf, r.csrf_quais = v_csll, '+'.join(quais)
        elif v_csll:                                   # sem grupo piscofins (código ausente): só a CSLL foi retida
            r.csll = v_csll
    total = local(raiz, 'vTotalRet')
    r.total_informado = para_float(total) if total is not None else None
    return r


def nota_de_sp(xml: str, papel: str = 'tomada') -> Nota:
    raiz = etree.fromstring(xml.encode('utf-8'))
    emissao = _data(local(raiz, 'DataEmissaoNFe'))
    aliquota = para_float(local(raiz, 'AliquotaServicos'))
    aliquota = aliquota * 100 if aliquota < 1 else aliquota              # SP manda 0.05 para 5%
    prest = grupo(raiz, 'CPFCNPJPrestador')
    toma = grupo(raiz, 'CPFCNPJTomador')
    end_prest = grupo(raiz, 'EnderecoPrestador')
    end_toma = grupo(raiz, 'EnderecoTomador')
    valor_iss = para_float(local(raiz, 'ValorISS'))
    iss_retido = (local(raiz, 'ISSRetido') or '').lower() == 'true'
    return Nota(
        fonte='SP', numero=local(raiz, 'NumeroNFe') or '',
        chave=f"SP-{local(raiz, 'NumeroNFe')}-{local(raiz, 'CodigoVerificacao')}",
        emissao=emissao, competencia=Competencia(emissao.year, emissao.month) if emissao else None,
        prestador_cnpj=digitos(local(prest, 'CNPJ')),
        prestador_nome=local(raiz, 'RazaoSocialPrestador') or '',
        prestador_municipio_ibge=local(end_prest, 'Cidade') or '3550308',
        tomador_cnpj=digitos(local(toma, 'CNPJ')),
        tomador_nome=local(raiz, 'RazaoSocialTomador') or '',
        tomador_municipio_ibge=local(end_toma, 'Cidade') or '',
        valor_servico=para_float(local(raiz, 'ValorServicos')),
        aliquota=aliquota, valor_iss=valor_iss,
        retencoes=retencoes_de_sp(raiz, emissao, valor_iss, iss_retido),     # é daqui que sai o iss_retido
        cancelada=(local(raiz, 'StatusNFe') or '') == 'C', papel=papel, xml=xml)


def nota_nacional(xml: str, cnpj_empresa: str = '') -> Nota:
    """Com `cnpj_empresa`, o papel sai do próprio XML: tomador → 'tomada', emitente → 'prestada', nenhum → ''."""
    raiz = etree.fromstring(xml.encode('utf-8'))
    emissao = _data(local(raiz, 'dhProc') or local(raiz, 'dhEmi'))       # competência = mês da EMISSÃO
    # tpEmit 1 = o prestador emitiu (infNFSe/emit é ele). 2/3 = tomador ou intermediário emitiram: o prestador está em
    # infDPS/prest. Hoje o ADN rejeita 2/3, mas o leiaute prevê — e ler emit nesse caso trocaria prestador por tomador.
    if (local(raiz, 'tpEmit') or '1') == '1':
        emit = grupo(raiz, 'emit', 'prest', 'prestador')
    else:
        emit = grupo(raiz, 'prest', 'prestador')
        if emit is None:                                                   # (elemento lxml não vale como booleano)
            emit = grupo(raiz, 'emit')
    toma = grupo(raiz, 'toma', 'tomador')
    chave = local(raiz, 'chNFSe') or next((e.get('Id') for e in raiz.iter()
                                           if re.fullmatch(r'NFS\d{47,53}', e.get('Id') or '')), '')
    tp_ret_iss = local(raiz, 'tpRetISSQN') or ''
    valor_iss = para_float(local(raiz, 'vISSQN'))
    vliq = local(raiz, 'vLiq')
    n = Nota(
        fonte='Nacional', numero=local(raiz, 'nNFSe') or '', chave=chave or '',
        emissao=emissao, competencia=Competencia(emissao.year, emissao.month) if emissao else None,
        prestador_cnpj=digitos(local(emit, 'CNPJ')),
        prestador_nome=local(emit, 'xNome') or '',
        prestador_municipio_ibge=local(emit, 'cMun') or local(raiz, 'cLocEmi') or '',
        tomador_cnpj=digitos(local(toma, 'CNPJ')),
        tomador_nome=local(toma, 'xNome') or '',
        tomador_municipio_ibge=local(toma, 'cMun') or '',
        valor_servico=para_float(local(raiz, 'vServ')),
        aliquota=para_float(local(raiz, 'pAliqAplic') or local(raiz, 'pAliq')),
        valor_iss=valor_iss,
        retencoes=retencoes_nacional(raiz, valor_iss, tp_ret_iss),           # é daqui que sai o iss_retido
        valor_liquido=para_float(vliq) if vliq is not None else None,
        xml=xml)
    if cnpj_empresa:
        n.papel = papel_de(n, cnpj_empresa)
    return n


def papel_de(nota: Nota, cnpj_empresa: str) -> str:
    """'tomada' se a empresa é o tomador, 'prestada' se é o prestador, '' se ela não aparece (intermediária, p. ex.)."""
    cnpj = digitos(cnpj_empresa)
    if nota.tomador_cnpj == cnpj:
        return 'tomada'
    if nota.prestador_cnpj == cnpj:
        return 'prestada'
    return ''


def chaves_canceladas(docs: list[dict]) -> set[str]:
    """Os EVENTOS de cancelamento chegam pela mesma distribuição: guardam a chave da nota cancelada."""
    canceladas = set()
    for d in docs:
        if 'EVENTO' in d['tipo'].upper() or '<infPedReg' in d['xml'] or 'e101101' in d['xml']:
            with contextlib.suppress(Exception):
                ch = local(etree.fromstring(d['xml'].encode('utf-8')), 'chNFSe')
                if ch:
                    canceladas.add(ch)
    return canceladas


# =====================================================================================================================
# 7. PASTAS E ESTADO
# =====================================================================================================================
RE_PASTA_MES = re.compile(r'^\s*(\d{1,2})\s*[-._ ]?\s*(\d{4})\s*$')


def pasta_da_competencia(pasta_cliente: Path, comp: Competencia) -> Path:
    """Acha-ou-cria NFE/<AAAA>/<MM-AAAA>. Reaproveita a pasta do mês em QUALQUER formato que a equipe tenha criado à
    mão ("09-2026", "09.2026", "092026", "09 2026") — o robô não cria uma segunda pasta para o mesmo mês."""
    base = pasta_cliente / 'NFE' / str(comp.ano)
    base.mkdir(parents=True, exist_ok=True)
    for sub in base.iterdir():
        m = RE_PASTA_MES.match(sub.name)
        if sub.is_dir() and m and (int(m.group(1)), int(m.group(2))) == (comp.mes, comp.ano):
            return sub
    alvo = base / comp.pasta
    alvo.mkdir(exist_ok=True)
    return alvo


class EstadoRobo:
    """Estado que sobrevive entre execuções: o último NSU por CNPJ e o 'ledger' do que o robô já escriturou."""

    def __init__(self, pasta: Path):
        self.pasta = pasta
        pasta.mkdir(parents=True, exist_ok=True)
        self.arquivo = pasta / 'estado.json'
        self.dados = json.loads(self.arquivo.read_text('utf-8')) if self.arquivo.exists() else {'nsu': {}, 'ledger': {}}

    def gravar(self):
        tmp = self.arquivo.with_suffix('.tmp')
        tmp.write_text(json.dumps(self.dados, ensure_ascii=False, indent=1), 'utf-8')
        tmp.replace(self.arquivo)                                       # gravação atômica

    def ultimo_nsu(self, cnpj: str) -> int:
        return int(self.dados['nsu'].get(cnpj, 0))

    def guardar_nsu(self, cnpj: str, nsu: int):
        self.dados['nsu'][cnpj] = max(nsu, self.ultimo_nsu(cnpj))
        self.gravar()

    def ja_escriturada(self, cnpj: str, chave: str) -> bool:
        return chave in self.dados['ledger'].get(cnpj, [])

    def registrar_escriturada(self, cnpj: str, chave: str):
        self.dados['ledger'].setdefault(cnpj, []).append(chave)
        self.gravar()

    def canceladas(self, cnpj: str) -> set[str]:
        """Chaves canceladas já vistas na distribuição deste CNPJ (em qualquer execução)."""
        return set(self.dados.get('canceladas', {}).get(cnpj, []))

    def registrar_canceladas(self, cnpj: str, chaves: Iterable[str]):
        """O evento de cancelamento só vem UMA vez pelo NSU. Se a nota estava arquivada (outra competência, ou papel
        não pedido naquela execução), ela só será lida depois — e precisa continuar cancelada. Por isso fica aqui."""
        novas = set(chaves) - self.canceladas(cnpj)
        if not novas:
            return
        self.dados.setdefault('canceladas', {}).setdefault(cnpj, []).extend(sorted(novas))
        self.gravar()

    def arquivar(self, cnpj: str, nota: Nota):
        """Documento de OUTRA competência, ou de um papel que não foi pedido nesta execução (o NSU traz tudo o que é
        novo): guardado, para quando aquela competência ou aquele papel for processado — nada se perde."""
        if not nota.competencia:
            return
        destino = self.pasta / 'arquivo' / cnpj / nota.competencia.pasta
        destino.mkdir(parents=True, exist_ok=True)
        (destino / f'{nome_seguro(nota.chave or nota.numero)}.xml').write_text(nota.xml, 'utf-8')

    def arquivadas(self, cnpj: str, comp: Competencia) -> list[str]:
        pasta = self.pasta / 'arquivo' / cnpj / comp.pasta
        return [p.read_text('utf-8') for p in sorted(pasta.glob('*.xml'))] if pasta.exists() else []


# =====================================================================================================================
# 8. ESPELHO EM PDF (opcional — o web service de SP não devolve PDF; o espelho reúne o que importa ao fiscal)
# =====================================================================================================================
def gerar_espelho_pdf(nota: Nota, destino: Path) -> bool:
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
    except ImportError:
        return False                                                  # espelho é conveniência: sem reportlab, segue
    r = nota.retencoes
    c = canvas.Canvas(str(destino), pagesize=A4)
    y = A4[1] - 60
    linhas = [('NFS-e', f'{nota.numero} ({nota.fonte}) — {nota.papel}'), ('Chave', nota.chave),
              ('Emissão', nota.emissao.strftime('%d/%m/%Y') if nota.emissao else ''),
              ('Prestador', f'{nota.prestador_nome} — {nota.prestador_cnpj}'),
              ('Município do prestador (IBGE)', nota.prestador_municipio_ibge),
              ('Tomador', f'{nota.tomador_nome} — {nota.tomador_cnpj}'.strip(' —')),
              ('Valor do serviço', f'R$ {nota.valor_servico:,.2f}'),
              ('Alíquota', f'{nota.aliquota:.2f}%'), ('ISS', f'R$ {nota.valor_iss:,.2f}'),
              ('ISS retido', f'sim, pelo {r.iss_por}' if nota.iss_retido else 'não')]
    for rotulo, valor in (('INSS retido', r.inss), ('IRRF retido', r.irrf), ('PIS retido', r.pis),
                          ('COFINS retido', r.cofins), ('CSLL retida', r.csll),
                          (f'CSRF retida ({r.csrf_quais})', r.csrf)):
        if valor:
            linhas.append((rotulo, f'R$ {valor:,.2f}'))
    linhas += [('Total retido', f'R$ {r.total:,.2f}'), ('Valor líquido', f'R$ {nota.valor_liquido:,.2f}'),
               ('Situação', 'CANCELADA' if nota.cancelada else 'autorizada')]
    if r.divergencia():
        linhas.append(('Conferir', r.divergencia()))
    for rotulo, valor in linhas:
        c.setFont('Helvetica-Bold', 10)
        c.drawString(50, y, rotulo)
        c.setFont('Helvetica', 10)
        c.drawString(230, y, str(valor)[:70])
        y -= 18
    c.save()
    return True


# =====================================================================================================================
# 9. CLASSIFICAÇÃO
# =====================================================================================================================
def deduplicar(notas: list[Nota]) -> list[Nota]:
    """A mesma nota pode vir pelas duas fontes (SP e Nacional): fica uma, de preferência a do Nacional."""
    vistas, saida = set(), []
    for n in sorted(notas, key=lambda n: n.fonte != 'Nacional'):         # a do Nacional tem preferência
        if n.identidade in vistas:
            continue
        vistas.add(n.identidade)
        saida.append(n)
    return saida


def classificar(notas: list[Nota], municipio_ibge: str) -> dict[str, list[Nota]]:
    """Para as TOMADAS: 'do_municipio' (já nasceu no sistema da prefeitura: não se declara), 'fora' (precisa ser
    escriturada) e 'canceladas'. Deduplica entre SP e Nacional pela identidade da nota (CNPJ do prestador, nº, valor)."""
    saida = {'do_municipio': [], 'fora': [], 'canceladas': []}
    for n in deduplicar(notas):
        if n.cancelada:
            saida['canceladas'].append(n)
        elif n.prestador_municipio_ibge == municipio_ibge:
            saida['do_municipio'].append(n)
        else:
            saida['fora'].append(n)
    return saida


# =====================================================================================================================
# 10. ESCRITURAÇÃO — interface do portal municipal + as travas de segurança (que valem para qualquer portal)
# =====================================================================================================================
@dataclasses.dataclass
class Previa:
    """O formulário PREENCHIDO no portal, ainda NÃO gravado — é aqui que as travas olham."""
    numero_documento: str                    # o nº que está no formulário (tem de ser o da nota)
    iss_calculado: float                     # o ISS que o portal calculou (tem de bater com valor × alíquota)


class SessaoPerdida(Exception):
    """O portal derrubou a sessão — o orquestrador renova e tenta de novo a MESMA nota."""


class PortalMunicipal(abc.ABC):
    """O que o robô precisa de um portal de escrituração. No robô real a implementação conversa com um portal JSF:
    primeiro por navegador (Selenium + Chrome DevTools), depois direto por requisições HTTP — mais rápido e estável.
    A autenticação é do ambiente de cada escritório e não faz parte desta versão."""

    @abc.abstractmethod
    def abrir_sessao(self, empresa: Empresa) -> None: ...

    @abc.abstractmethod
    def escrituradas(self, empresa: Empresa, comp: Competencia) -> set[tuple]:
        """Identidades (CNPJ do prestador, nº, valor) já declaradas no mês — a consulta anti-duplicidade."""

    @abc.abstractmethod
    def preencher(self, empresa: Empresa, nota: Nota) -> Previa:
        """Preenche o formulário de declaração da nota e devolve o que o portal mostra, SEM gravar."""

    @abc.abstractmethod
    def gravar(self) -> bytes:
        """Confirma o formulário preenchido e devolve o comprovante (PDF)."""

    @abc.abstractmethod
    def descartar(self) -> None:
        """Abandona o formulário preenchido (a trava não deixou gravar)."""

    @abc.abstractmethod
    def emitir_guia_iss_retido(self, empresa: Empresa, comp: Competencia) -> Optional[bytes]: ...

    @abc.abstractmethod
    def exportar_xml_tomados(self, empresa: Empresa, comp: Competencia) -> tuple[list[str], dict[str, str]]:
        """XML exportados pelo portal + o mapa {sequencial do portal: nº real do documento} (seção 11)."""


def conferir_antes_de_gravar(nota: Nota, previa: Previa, tolerancia: float) -> Optional[str]:
    """As TRAVAS, antes de gravar. Devolve o motivo para NÃO gravar, ou None. Evita, por exemplo, o ISS errado de uma
    alíquota padrão do portal que não foi trocada pela da nota."""
    if previa.numero_documento.lstrip('0') != nota.numero.lstrip('0'):
        return f'nº do documento no portal ({previa.numero_documento}) ≠ nº da nota ({nota.numero})'
    esperado = round(nota.valor_servico * nota.aliquota / 100, 2)
    if abs(previa.iss_calculado - esperado) > tolerancia:
        return f'ISS do portal R$ {previa.iss_calculado:.2f} ≠ valor × alíquota R$ {esperado:.2f}'
    return None


def escriturar_empresa(portal: PortalMunicipal, empresa: Empresa, comp: Competencia, fora: list[Nota],
                       pasta_comp: Path, estado: EstadoRobo, cfg: Config, res: ResultadoEmpresa,
                       tentativas_sessao: int = 3) -> None:
    portal.abrir_sessao(empresa)
    ja = portal.escrituradas(empresa, comp)                             # 1) ANTI-DUPLICIDADE
    pasta_pdf = pasta_comp / 'Tomados' / 'Notas Escrituradas'
    for nota in fora:
        if nota.identidade in ja or estado.ja_escriturada(empresa.cnpj, nota.chave):
            res.ja_escrituradas += 1
            res.notas.append((nota, 'já escriturada'))
            continue
        motivo, pdf = None, None
        for tentativa in range(1, tentativas_sessao + 1):
            try:
                previa = portal.preencher(empresa, nota)                # 2) PREENCHE (sem gravar)
                motivo = conferir_antes_de_gravar(nota, previa, cfg.tolerancia_iss)   # 3) TRAVAS
                if motivo:
                    portal.descartar()
                else:
                    pdf = portal.gravar()                               # 4) GRAVA
                break
            except SessaoPerdida:
                log.warning('   sessão caiu (tentativa %d) — renovando e refazendo a nota %s', tentativa, nota.numero)
                portal.abrir_sessao(empresa)
        else:
            res.erros.append(f'nota {nota.numero}: o portal não manteve a sessão')
            res.notas.append((nota, 'erro: sessão'))
            continue
        if motivo:
            res.erros.append(f'nota {nota.numero}: {motivo}')
            res.notas.append((nota, f'NÃO escriturada: {motivo}'))
            continue
        if pdf:                                                         # 5) COMPROVANTE
            pasta_pdf.mkdir(parents=True, exist_ok=True)
            (pasta_pdf / f'{nome_seguro(nota.prestador_nome, 40)} NF {nota.numero}.pdf').write_bytes(pdf)
        estado.registrar_escriturada(empresa.cnpj, nota.chave)
        res.escrituradas += 1
        res.notas.append((nota, 'escriturada'))
    confirmadas = portal.escrituradas(empresa, comp)                    # 6) CONFIRMA no portal
    faltam = [n for n, s in res.notas if s == 'escriturada' and n.identidade not in confirmadas]
    if faltam:
        res.erros.append(f'{len(faltam)} nota(s) gravada(s) que o portal ainda não lista — conferir')
    # 7) GUIA do ISS retido — só o que a própria empresa (tomadora) reteve; retido pelo intermediário não é dela
    if any(n.iss_retido and n.retencoes.iss_por != 'intermediário'
           for n, s in res.notas if s in ('escriturada', 'já escriturada')):
        guia = portal.emitir_guia_iss_retido(empresa, comp)
        if guia:
            (pasta_comp / 'Tomados' / f'Guia ISS retido {comp.pasta}.pdf').write_bytes(guia)
            res.guia_emitida = True


# =====================================================================================================================
# 11. EXPORTAÇÃO PARA O SISTEMA CONTÁBIL
# =====================================================================================================================
def renumerar_xml(xml: str, mapa: dict[str, str]) -> str:
    """Na exportação, o portal troca o nº da nota por um SEQUENCIAL interno; o sistema contábil precisa do nº real do
    documento. O mapa sai da relação de tomados do próprio portal (sequencial → nº do documento escriturado)."""
    raiz = etree.fromstring(xml.encode('utf-8'))
    for el in raiz.iter():
        if isinstance(el.tag, str) and etree.QName(el).localname in ('Numero', 'NumeroNfse', 'nNFSe'):
            real = mapa.get((el.text or '').strip())
            if real:
                el.text = real
    return etree.tostring(raiz, xml_declaration=True, encoding='UTF-8').decode('utf-8')


def exportar_para_sistema_contabil(portal: PortalMunicipal, empresa: Empresa, comp: Competencia, pasta_comp: Path) -> int:
    xmls, mapa = portal.exportar_xml_tomados(empresa, comp)
    destino = pasta_comp / 'Tomados' / 'Importação XML Tomados'
    destino.mkdir(parents=True, exist_ok=True)
    for i, xml in enumerate(xmls, 1):
        (destino / f'tomados_{comp.pasta}_{i:03d}.xml').write_text(renumerar_xml(xml, mapa), 'utf-8')
    return len(xmls)


# =====================================================================================================================
# 12. RELATÓRIO E E-MAIL
# =====================================================================================================================
COLUNAS_EMPRESAS = ['Código', 'Empresa', 'CNPJ', 'Extraídas', 'Tomadas', 'Prestadas', 'Fora do município',
                    'Escrituradas agora', 'Já escrituradas', 'Guia ISS', 'Retido nas tomadas', 'Retido nas prestadas',
                    'Pendências']
COLUNAS_NOTAS = ['Código', 'Papel', 'Fonte', 'Nº', 'Emissão', 'Prestador', 'CNPJ prestador', 'Tomador', 'CNPJ tomador',
                 'Município (IBGE)', 'Valor', 'Alíquota %', 'ISS', 'ISS retido', 'INSS retido', 'IRRF retido',
                 'PIS retido', 'COFINS retido', 'CSLL retida', 'CSRF retida', 'Total retido', 'Valor líquido',
                 'Retenções', 'Situação']
COLUNAS_RETENCOES = ['Código', 'Empresa', 'Papel', 'Notas com retenção', 'ISS', 'ISS pelo intermediário', 'INSS', 'IRRF',
                     'PIS', 'COFINS', 'CSLL', 'CSRF (PIS+COFINS+CSLL)', 'Total']


def gerar_relatorio(resultados: list[ResultadoEmpresa], comp: Competencia, destino: Path) -> Path:
    """Excel com três abas — Empresas (o resumo), Notas (o detalhe, com cada retenção em sua coluna) e Retenções (o
    total por empresa, papel e tributo: o que a empresa recolhe das tomadas e abate das prestadas)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    emp = wb.active
    emp.title = 'Empresas'
    emp.append(COLUNAS_EMPRESAS)
    notas = wb.create_sheet('Notas')
    notas.append(COLUNAS_NOTAS)
    ret = wb.create_sheet('Retenções')
    ret.append(COLUNAS_RETENCOES)
    for r in resultados:
        e = r.empresa
        emp.append([e.codigo, e.razao, e.cnpj, r.extraidas, r.tomadas, r.prestadas, r.fora_do_municipio,
                    r.escrituradas, r.ja_escrituradas, 'sim' if r.guia_emitida else '',
                    r.retido('tomada'), r.retido('prestada'), ' | '.join(r.erros)])
        for n, situacao in r.notas:
            t = n.retencoes
            notas.append([e.codigo, n.papel, n.fonte, n.numero, n.emissao.strftime('%d/%m/%Y') if n.emissao else '',
                          n.prestador_nome, n.prestador_cnpj, n.tomador_nome, n.tomador_cnpj,
                          n.prestador_municipio_ibge, n.valor_servico, n.aliquota, n.valor_iss,
                          'sim' if n.iss_retido else 'não', t.inss, t.irrf, t.pis, t.cofins, t.csll, t.csrf,
                          t.total, n.valor_liquido, t.descricao(), situacao])
        for papel in PAPEIS:
            # "Total" é o que a empresa recolhe (tomadas: sem o ISS do intermediário) ou abate (prestadas: tudo)
            com = [n for n, _ in r.notas if n.papel == papel and not n.cancelada and n.retencoes.tem]
            if com:
                soma = lambda campo: round(sum(getattr(n.retencoes, campo) for n in com), 2)   # noqa: E731
                iss_interm = round(sum(n.retencoes.iss for n in com if n.retencoes.iss_por == 'intermediário'), 2)
                ret.append([e.codigo, e.razao, papel, len(com), round(soma('iss') - iss_interm, 2), iss_interm,
                            soma('inss'), soma('irrf'), soma('pis'), soma('cofins'), soma('csll'), soma('csrf'),
                            round(sum(retido_pela_empresa(n) for n in com), 2)])
    for ws in (emp, notas, ret):
        for c in ws[1]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor='1F4E5F')
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(48, max(10, *(len(str(c.value or '')) for c in col)))
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino


def enviar_por_email(cfg: Config, para: str, anexo: Path, comp: Competencia) -> bool:
    """SMTP com STARTTLS. Servidor e usuário vêm da configuração; a senha, do ambiente (NFSE_SMTP_SENHA)."""
    senha = os.environ.get('NFSE_SMTP_SENHA', '')
    if not (cfg.smtp_host and cfg.smtp_usuario and senha and para):
        log.info('e-mail não configurado — relatório só salvo em %s', anexo)
        return False
    msg = EmailMessage()
    msg['Subject'] = f'NFS-e tomadas e prestadas — {comp}'
    msg['From'], msg['To'] = cfg.smtp_usuario, para
    msg.set_content(f'Segue o relatório de NFS-e (tomadas, prestadas e retenções) da competência {comp}.')
    msg.add_attachment(anexo.read_bytes(), maintype='application',
                       subtype='vnd.openxmlformats-officedocument.spreadsheetml.sheet', filename=anexo.name)
    with smtplib.SMTP(cfg.smtp_host, cfg.smtp_porta, timeout=60) as s:
        s.starttls()
        s.login(cfg.smtp_usuario, senha)
        s.send_message(msg)
    return True


# =====================================================================================================================
# 13. ORQUESTRADOR
# =====================================================================================================================
class Progresso:
    """Andamento com estimativa de término — no robô real alimenta um painel de monitoramento ao vivo."""

    def __init__(self, total: int):
        self.total, self.inicio, self.feitos = total, time.monotonic(), 0

    def avanca(self, rotulo: str):
        self.feitos += 1
        decorrido = time.monotonic() - self.inicio
        resta = decorrido / self.feitos * (self.total - self.feitos)
        log.info('[%d/%d] %s  (faltam ~%ds)', self.feitos, self.total, rotulo, resta)


def extrair_empresa(empresa: Empresa, comp: Competencia, cert: Optional[CertificadoA1], senhas: SenhaProvider,
                    estado: EstadoRobo, fontes: dict, papeis: Iterable[str] = PAPEIS) -> list[Nota]:
    """ETAPA 1+2 de uma empresa: busca nas duas fontes as notas que ela TOMOU e/ou PRESTOU (conforme `papeis`), fica
    com as da competência pedida; o resto do Nacional vai para o arquivo. Usa também o que já tinha sido arquivado
    em execuções anteriores."""
    papeis = [p for p in PAPEIS if p in set(papeis)]                    # ordem fixa: tomada, depois prestada
    notas: list[Nota] = []
    # cancelamentos: os desta distribuição + os já vistos antes (o evento vem uma vez só; a nota pode ter ficado
    # arquivada e só ser lida agora)
    canceladas: set[str] = estado.canceladas(empresa.cnpj)
    if cert is None:
        raise RuntimeError('sem certificado A1 válido')
    senha = senhas.senha(empresa.cnpj)
    if 'SP' in fontes:
        sp = fontes['SP']
        if 'tomada' in papeis:
            notas += [nota_de_sp(x, 'tomada') for x in sp.consultar_recebidas(cert, senha, comp)]
        if 'prestada' in papeis:
            notas += [nota_de_sp(x, 'prestada') for x in sp.consultar_emitidas(cert, senha, comp)]
    if 'Nacional' in fontes:
        docs, ultimo = fontes['Nacional'].distribuir(cert, senha, estado.ultimo_nsu(empresa.cnpj))
        novas = chaves_canceladas(docs)
        estado.registrar_canceladas(empresa.cnpj, novas)                 # antes de avançar o NSU
        canceladas |= novas
        for d in docs:
            with contextlib.suppress(Exception):
                n = nota_nacional(d['xml'], empresa.cnpj)
                if not n.papel or not n.numero:
                    continue                                          # intermediária, ou evento
                n.cancelada = n.chave in canceladas
                if n.competencia == comp and n.papel in papeis:
                    notas.append(n)
                else:                                                 # outra competência ou papel não pedido agora
                    estado.arquivar(empresa.cnpj, n)
        estado.guardar_nsu(empresa.cnpj, ultimo)                       # só depois de tudo salvo
    for x in estado.arquivadas(empresa.cnpj, comp):
        with contextlib.suppress(Exception):
            n = nota_nacional(x, empresa.cnpj)
            if n.papel in papeis and n.numero:
                n.cancelada = n.chave in canceladas
                notas.append(n)
    # O Nacional é incremental: o que veio numa execução anterior desta competência já está salvo na pasta da
    # empresa e não vem de novo. Entra aqui, para o relatório ficar completo em qualquer rodada (e cancelamento
    # que chegou agora vale para a nota antiga).
    if empresa.pasta:
        pasta_comp = pasta_da_competencia(empresa.pasta, comp)
        vistas = {n.chave for n in notas}
        for papel in papeis:
            raiz = PASTA_DO_PAPEL[papel]
            pasta_nacional = pasta_comp / raiz / f'{raiz} Portal Nacional'
            for arquivo in (sorted(pasta_nacional.rglob('*.xml')) if pasta_nacional.is_dir() else []):
                with contextlib.suppress(Exception):
                    n = nota_nacional(arquivo.read_text('utf-8'), empresa.cnpj)
                    if n.chave and n.chave not in vistas and n.papel == papel and n.competencia == comp:
                        n.cancelada = n.chave in canceladas or 'Canceladas' in arquivo.parts
                        notas.append(n)
                        vistas.add(n.chave)
    return notas


def salvar_xmls(notas: list[Nota], pasta_comp: Path) -> None:
    """Tomados/Tomados SP|Tomados Portal Nacional/Autorizadas|Canceladas e o mesmo em Prestados/. O arquivo leva o
    nome de quem está do outro lado (o prestador numa tomada, o tomador numa prestada) e o nº da nota."""
    for n in notas:
        raiz = PASTA_DO_PAPEL.get(n.papel, 'Tomados')
        sub = pasta_comp / raiz / (f'{raiz} SP' if n.fonte == 'SP' else f'{raiz} Portal Nacional') \
            / ('Canceladas' if n.cancelada else 'Autorizadas')
        sub.mkdir(parents=True, exist_ok=True)
        quem = n.contraparte or (n.tomador_cnpj if n.papel == 'prestada' else n.prestador_cnpj)
        base = sub / f'{nome_seguro(quem, 40)} NF {n.numero}'
        base.with_suffix('.xml').write_text(n.xml, 'utf-8')
        gerar_espelho_pdf(n, base.with_suffix('.pdf'))


def situacao_da_prestada(n: Nota) -> str:
    if n.cancelada:
        return 'prestada cancelada'
    return 'prestada — com retenção' if n.retencoes.tem else 'prestada'


def executar(cfg: Config, comp: Competencia, empresas: list[Empresa], senhas: SenhaProvider,
             certificados: dict[str, CertificadoA1], portal: Optional[PortalMunicipal], fontes: dict,
             extrator: Optional[Callable[[Empresa], list[Nota]]] = None,
             resultados: Optional[list] = None) -> list[Path]:
    """Roda o processo inteiro para a lista de empresas e devolve os relatórios gerados (um por responsável).
    Quem quiser o detalhe (a interface local, por exemplo) passa uma lista em `resultados`: ela é preenchida
    com um ResultadoEmpresa por empresa."""
    estado = EstadoRobo(cfg.pasta_estado)
    progresso = Progresso(len(empresas))
    if resultados is None:
        resultados = []
    for i, empresa in enumerate(empresas):
        if i:
            time.sleep(cfg.pausa_entre_empresas)
        res = ResultadoEmpresa(empresa)
        resultados.append(res)
        try:
            empresa.pasta = empresa.pasta or cfg.pasta_saida / nome_seguro(f'{empresa.codigo} - {empresa.razao}')
            pasta_comp = pasta_da_competencia(empresa.pasta, comp)
            notas = extrator(empresa) if extrator else extrair_empresa(
                empresa, comp, certificados.get(empresa.cnpj), senhas, estado, fontes, cfg.papeis)
            res.extraidas = len(notas)
            salvar_xmls(notas, pasta_comp)
            # TOMADAS: do município × fora × canceladas; as de fora vão para o portal
            grupos = classificar([n for n in notas if n.papel != 'prestada'], cfg.municipio_escrituracao_ibge)
            res.tomadas = sum(len(v) for v in grupos.values())
            res.fora_do_municipio = len(grupos['fora'])
            for n in grupos['do_municipio']:
                res.notas.append((n, 'do município (não se declara)'))
            for n in grupos['canceladas']:
                res.notas.append((n, 'cancelada'))
            # a escrituração é no portal do município DA EMPRESA: só vale para quem é desse município
            if portal and empresa.municipio_ibge == cfg.municipio_escrituracao_ibge and grupos['fora']:
                escriturar_empresa(portal, empresa, comp, grupos['fora'], pasta_comp, estado, cfg, res)
                exportar_para_sistema_contabil(portal, empresa, comp, pasta_comp)
            else:
                for n in grupos['fora']:
                    res.notas.append((n, 'fora do município — escriturar no portal do tomador'))
            # PRESTADAS: nada a escriturar — o que importa é o que o cliente reteve (vai para o relatório)
            prestadas = deduplicar([n for n in notas if n.papel == 'prestada'])
            res.prestadas = len(prestadas)
            for n in prestadas:
                res.notas.append((n, situacao_da_prestada(n)))
        except Exception as e:                                         # falha isolada: a próxima empresa segue
            log.exception('   %s: %s', empresa.razao, e)
            res.erros.append(f'{type(e).__name__}: {e}')
        progresso.avanca(f'{empresa.codigo} {empresa.razao}: {res.tomadas} tomada(s), {res.prestadas} prestada(s), '
                         f'{res.escrituradas} escriturada(s)'
                         + (f', {res.ja_escrituradas} já escriturada(s)' if res.ja_escrituradas else '')
                         + (f', R$ {res.retido("prestada"):.2f} retido(s) nas prestadas' if res.retido('prestada') else '')
                         + (f', {len(res.erros)} pendência(s)' if res.erros else ''))
    relatorios = []
    for resp in sorted({r.empresa.responsavel for r in resultados}):
        do_resp = [r for r in resultados if r.empresa.responsavel == resp]
        arq = gerar_relatorio(do_resp, comp, cfg.pasta_saida / '_relatorios' / f'{nome_seguro(resp)} {comp.pasta}.xlsx')
        relatorios.append(arq)
    return relatorios


# =====================================================================================================================
# 14. DEMONSTRAÇÃO E LINHA DE COMANDO
# =====================================================================================================================
def cnpj_ficticio(rnd: random.Random) -> str:
    """CNPJ com dígitos verificadores válidos, gerado ao acaso — só para a demonstração."""
    base = [rnd.randint(0, 9) for _ in range(8)] + [0, 0, 0, 1]
    for pesos in ([5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2], [6, 5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]):
        resto = sum(d * p for d, p in zip(base, pesos)) % 11
        base.append(0 if resto < 2 else 11 - resto)
    return ''.join(map(str, base))


class PortalSimulado(PortalMunicipal):
    """Portal em memória, com o comportamento que importa: guarda o que já foi declarado, calcula o ISS pela alíquota
    informada, às vezes derruba a sessão — e tem um "defeito" de propósito (alíquota padrão de 5% em 1 de cada 7
    notas) para mostrar a trava de conferência do ISS funcionando."""

    def __init__(self, rnd: random.Random):
        self.rnd, self.base, self.contador, self.formulario = rnd, {}, 0, None

    def abrir_sessao(self, empresa):
        self.formulario = None

    def escrituradas(self, empresa, comp):
        return {n.identidade for n in self.base.get((empresa.cnpj, comp), [])}

    def preencher(self, empresa, nota):
        self.contador += 1
        if self.rnd.random() < 0.08:
            raise SessaoPerdida()
        aliquota = 5.0 if self.contador % 7 == 0 else nota.aliquota      # o "defeito" de propósito
        self.formulario = (empresa, nota)
        return Previa(nota.numero, round(nota.valor_servico * aliquota / 100, 2))

    def gravar(self):
        empresa, nota = self.formulario
        self.base.setdefault((empresa.cnpj, nota.competencia), []).append(nota)
        self.formulario = None
        return b'%PDF-1.4\n% comprovante simulado\n'

    def descartar(self):
        self.formulario = None

    def emitir_guia_iss_retido(self, empresa, comp):
        return b'%PDF-1.4\n% guia simulada\n'

    def exportar_xml_tomados(self, empresa, comp):
        xmls, mapa = [], {}
        for i, n in enumerate(self.base.get((empresa.cnpj, comp), []), 1):
            xmls.append(f'<CompNfse><Nfse><InfNfse><Numero>{i}</Numero></InfNfse></Nfse></CompNfse>')
            mapa[str(i)] = n.numero
        return xmls, mapa


def retencoes_ficticias(rnd: random.Random, valor: float) -> Retencoes:
    """Para a demonstração: INSS (11%), IRRF (1,5%) e a soma CSRF (4,65%), cada um presente ou não, ao acaso."""
    r = Retencoes()
    if rnd.random() < 0.35:
        r.inss = round(valor * 0.11, 2)
    if rnd.random() < 0.45:
        r.irrf = round(valor * 0.015, 2)
    if rnd.random() < 0.45:
        r.csrf, r.csrf_quais = round(valor * 0.0465, 2), 'PIS+COFINS+CSLL'
    return r


def demonstracao(saida: Path, resultados: Optional[list] = None) -> list[Path]:
    """Fluxo completo, offline, com empresas e notas FICTÍCIAS geradas aqui mesmo."""
    rnd = random.Random(42)
    comp = Competencia.anterior_a(dt.date.today())
    cfg = Config(pasta_saida=saida, pasta_estado=saida / '_estado', pausa_entre_empresas=0)
    municipio = cfg.municipio_escrituracao_ibge
    outros = ['3550308', '3304557', '4106902', '3106200']             # outros municípios (códigos IBGE públicos)
    empresas = [Empresa(str(100 + i), f'Empresa Exemplo {chr(65 + i)} Ltda', cnpj_ficticio(rnd),
                        municipio if i < 4 else rnd.choice(outros), ['Analista 1', 'Analista 2'][i % 2])
                for i in range(6)]

    def extrator(e: Empresa) -> list[Nota]:
        rnd = random.Random(f'notas-{e.codigo}')                       # as MESMAS notas a cada execução
        notas = []
        for k in range(rnd.randint(3, 9)):                              # o que a empresa TOMOU
            valor = round(rnd.uniform(300, 15000), 2)
            aliq = rnd.choice([2.0, 2.5, 3.0, 5.0])
            emissao = comp.inicio + dt.timedelta(days=rnd.randint(0, 27))
            n = Nota(fonte=rnd.choice(['SP', 'Nacional']), numero=str(rnd.randint(1000, 99999)), chave=f'DEMO{e.codigo}{k}',
                     emissao=emissao, competencia=comp, prestador_cnpj=cnpj_ficticio(rnd),
                     prestador_nome=f'Prestador Exemplo {k + 1}',
                     prestador_municipio_ibge=rnd.choice([municipio] + outros * 2), tomador_cnpj=e.cnpj,
                     tomador_nome=e.razao, tomador_municipio_ibge=e.municipio_ibge,
                     valor_servico=valor, aliquota=aliq, valor_iss=round(valor * aliq / 100, 2),
                     iss_retido=rnd.random() < 0.3, cancelada=rnd.random() < 0.08,
                     xml=f'<NFSe><nNFSe>{k}</nNFSe><demo>fictícia</demo></NFSe>')
            notas.append(n)
        ret = random.Random(f'retencoes-{e.codigo}')                    # retenções federais nas tomadas
        for n in notas:
            r = retencoes_ficticias(ret, n.valor_servico)
            r.iss, r.iss_por = n.retencoes.iss, n.retencoes.iss_por
            n.retencoes = r                                              # o líquido acompanha sozinho (propriedade)
        if notas:                                                       # a mesma nota pelas duas fontes (dedup),
            notas.append(dataclasses.replace(notas[0], fonte='SP' if notas[0].fonte == 'Nacional' else 'Nacional'))
            notas[-1].retencoes = dataclasses.replace(notas[0].retencoes)   # com as MESMAS retenções da original
        pre = random.Random(f'prestadas-{e.codigo}')                    # o que a empresa PRESTOU
        for k in range(pre.randint(2, 6)):
            valor = round(pre.uniform(500, 20000), 2)
            aliq = pre.choice([2.0, 3.0, 5.0])
            notas.append(Nota(fonte=pre.choice(['SP', 'Nacional']), numero=str(pre.randint(1000, 99999)),
                              chave=f'DEMOP{e.codigo}{k}', emissao=comp.inicio + dt.timedelta(days=pre.randint(0, 27)),
                              competencia=comp, prestador_cnpj=e.cnpj, prestador_nome=e.razao,
                              prestador_municipio_ibge=e.municipio_ibge, tomador_cnpj=cnpj_ficticio(pre),
                              tomador_nome=f'Cliente Exemplo {k + 1}', tomador_municipio_ibge=pre.choice([municipio] + outros),
                              valor_servico=valor, aliquota=aliq, valor_iss=round(valor * aliq / 100, 2),
                              iss_retido=pre.random() < 0.3, retencoes=retencoes_ficticias(pre, valor),
                              cancelada=pre.random() < 0.05, papel='prestada',
                              xml=f'<NFSe><nNFSe>P{k}</nNFSe><demo>fictícia</demo></NFSe>'))
        return notas

    # rode duas vezes: na segunda, o que já foi escriturado aparece como "já escriturada" — nada é declarado em dobro
    return executar(cfg, comp, empresas, SenhaDoAmbiente(), {}, PortalSimulado(random.Random()), {},
                    extrator=extrator, resultados=resultados)


def carregar_empresas(arquivo: Path) -> list[Empresa]:
    """JSON: [{"codigo": "...", "razao": "...", "cnpj": "...", "municipio_ibge": "...", "responsavel": "..."}]"""
    return [Empresa(**{k: v for k, v in e.items() if k in {f.name for f in dataclasses.fields(Empresa)}})
            for e in json.loads(arquivo.read_text('utf-8'))]


def papeis_de(so: Optional[str]) -> tuple:
    """`--so tomadas` / `--so prestadas` → só aquele papel; sem a opção, os dois."""
    return {'tomadas': ('tomada',), 'prestadas': ('prestada',)}.get(so or '', PAPEIS)


def main(argv=None):
    p = argparse.ArgumentParser(description='Robô NFS-e — tomadas e prestadas: extração, retenções, organização e escrituração')
    p.add_argument('--demo', action='store_true', help='roda o fluxo completo com dados fictícios, offline')
    p.add_argument('--competencia', help='MM/AAAA (padrão: o mês anterior)')
    p.add_argument('--empresas', type=Path, help='JSON com as empresas (ver carregar_empresas)')
    p.add_argument('--certificados', type=Path, help='pasta dos certificados A1 (uma subpasta por cliente)')
    p.add_argument('--saida', type=Path, default=Path('saida_nfse'))
    p.add_argument('--responsavel', help='só as empresas deste responsável')
    p.add_argument('--so', choices=['tomadas', 'prestadas'], help='buscar só as notas tomadas ou só as prestadas (padrão: as duas)')
    p.add_argument('--interface', action='store_true',
                   help='abre a interface local no navegador (certificado e senha não saem desta máquina)')
    p.add_argument('--porta', type=int, default=None, help='porta da interface local (padrão: 8747)')
    p.add_argument('--sem-navegador', action='store_true', help='com --interface: sobe a página sem abrir o navegador')
    a = p.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format='%(asctime)s %(message)s', datefmt='%H:%M:%S')

    if a.interface:
        from interface.servidor import iniciar, PORTA_PADRAO
        iniciar(porta=a.porta or PORTA_PADRAO, abrir_navegador=not a.sem_navegador)
        return
    if a.demo:
        for r in demonstracao(a.saida):
            log.info('relatório: %s', r)
        log.info('demonstração concluída — veja as pastas e os relatórios em %s', a.saida.resolve())
        return
    if not (a.empresas and a.certificados):
        p.error('informe --empresas e --certificados (ou use --demo)')
    comp = Competencia.de_texto(a.competencia) if a.competencia else Competencia.anterior_a(dt.date.today())
    empresas = [e for e in carregar_empresas(a.empresas) if not a.responsavel or e.responsavel == a.responsavel]
    senhas = SenhaDoAmbiente()
    certs = certificados_por_cnpj(a.certificados, senhas, [e.cnpj for e in empresas])
    cfg = Config(pasta_saida=a.saida, pasta_estado=a.saida / '_estado', papeis=papeis_de(a.so))
    fontes = {'SP': ClienteSaoPaulo(), 'Nacional': ClienteNacional()}
    # sem um PortalMunicipal real configurado, o robô extrai, organiza, classifica e relata — e lista o que escriturar
    for r in executar(cfg, comp, empresas, senhas, certs, portal=None, fontes=fontes):
        log.info('relatório: %s', r)
        enviar_por_email(cfg, os.environ.get('NFSE_EMAIL_PARA', ''), r, comp)


if __name__ == '__main__':
    main()
